from flask import Flask, render_template, request, jsonify, send_file, session, redirect, url_for
# Import các module chức năng riêng biệt theo mô hình modular (Đã tích hợp assign_assets_handler)
from modules import topic_ideator, profile_manager, auto_extractor, research_script_writer, direct_script_writer, assign_assets_handler, get_prompt_image_handler, prescan_handler, asset_prompts, video_metadata_handler, thumbnail_handler, camera_handler, get_prompt_video_handler
# Module dùng chung để gọi AI (hỗ trợ cả ShopAIKey "sk-" và Gemini gốc)
from modules import ai_client
import os
import json
import re
import zipfile
import io
import uuid
import time
import secrets
from functools import wraps
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Nạp biến môi trường từ file .env (nếu có) cho lúc chạy local
try:
    from dotenv import load_dotenv
    load_dotenv()
except Exception:
    pass

# Khởi động ứng dụng Flask, tự động nhận diện thư mục 'templates' cùng cấp
app = Flask(__name__)

# --- CẤU HÌNH XÁC THỰC (AUTHENTICATION) ---
# Hai cơ chế song song:
#   1. Google OAuth (ưu tiên) — mỗi Gmail là một tài khoản độc lập, dữ liệu
#      project/profile tách riêng. Bật khi có GOOGLE_CLIENT_ID + SECRET.
#   2. Username/Password dùng chung — cơ chế cũ, dùng khi chưa cấu hình OAuth.
from modules import auth as auth_module
from authlib.integrations.flask_client import OAuth

AUTH_USERNAME = os.environ.get("APP_USERNAME", "")
AUTH_PASSWORD = os.environ.get("APP_PASSWORD", "")
# SECRET_KEY dùng để ký session cookie.
#
# QUAN TRỌNG khi chạy nhiều worker (gunicorn --workers 2):
# nếu mỗi worker tự sinh key ngẫu nhiên thì cookie do worker này ký,
# worker kia không đọc được -> người dùng bị đá ra mỗi lần chuyển trang.
# Vì vậy trên production BẮT BUỘC phải khai SECRET_KEY trong biến môi trường.
_secret = os.environ.get("SECRET_KEY", "").strip()
if not _secret:
    # Chạy local thì tự sinh cho tiện, nhưng phải cảnh báo rõ ràng
    _secret = secrets.token_hex(32)
    print(
        "[!] CANH BAO: Chua khai bien SECRET_KEY.\n"
        "    Moi worker se sinh key rieng -> dang nhap bi mat khi chuyen trang.\n"
        "    Tren Railway: Variables -> them SECRET_KEY = <chuoi ngau nhien dai>."
    )
app.secret_key = _secret

# Thời gian phiên đăng nhập duy trì (giây) - mặc định 7 ngày
app.permanent_session_lifetime = 60 * 60 * 24 * 7

# Cookie phiên: chống gửi kèm sang site khác, và ép HTTPS khi chạy production.
app.config.update(
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=bool(os.environ.get("RAILWAY_ENVIRONMENT")
                               or os.environ.get("SESSION_COOKIE_SECURE")),
)

# Railway đứng sau reverse proxy: không có ProxyFix thì Flask tưởng request là
# http, url_for(_external=True) sinh redirect_uri sai -> Google OAuth báo lỗi.
try:
    from werkzeug.middleware.proxy_fix import ProxyFix
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
except Exception:
    pass

# Những route KHÔNG cần đăng nhập (trang login/logout/tài nguyên tĩnh)
PUBLIC_PATHS = {"/login", "/logout", "/favicon.ico",
                "/login/google", "/auth/google/callback"}

# --- ĐĂNG KÝ GOOGLE OAUTH ---
oauth = OAuth(app)
if auth_module.google_oauth_enabled():
    oauth.register(
        name="google",
        client_id=auth_module.GOOGLE_CLIENT_ID,
        client_secret=auth_module.GOOGLE_CLIENT_SECRET,
        server_metadata_url=auth_module.GOOGLE_DISCOVERY_URL,
        client_kwargs={"scope": "openid email profile"},
    )


def auth_required() -> bool:
    """
    Hệ thống có bắt buộc đăng nhập hay không.

    Chạy local mà chưa khai gì cả -> không bắt buộc, vào thẳng cho tiện.
    """
    return auth_module.google_oauth_enabled() or bool(AUTH_USERNAME and AUTH_PASSWORD)


def login_required_html(view_fn):
    """Decorator giữ nguyên hành vi render trang HTML: nếu chưa login thì chuyển về /login."""
    @wraps(view_fn)
    def wrapper(*args, **kwargs):
        if not session.get("authenticated"):
            return redirect(url_for("login_page"))
        return view_fn(*args, **kwargs)
    return wrapper


@app.before_request
def require_auth():
    """Chặn MỌI request khi chưa đăng nhập.
    - Trang HTML: redirect tới /login
    - API (đường dẫn /api/...): trả về JSON 401
    - Nếu chưa cấu hình gì cả thì bỏ qua (tiện chạy local)
    """
    # Chưa cấu hình cơ chế đăng nhập nào → không bắt buộc (chỉ dùng lúc local)
    if not auth_required():
        return None

    path = request.path
    if path in PUBLIC_PATHS or path.startswith("/static"):
        return None
    if session.get("authenticated"):
        return None
    if path.startswith("/api/"):
        return jsonify({"status": "error", "message": "Chưa đăng nhập. Vui lòng đăng nhập để tiếp tục.", "code": "UNAUTHORIZED"}), 401
    return redirect(url_for("login_page"))


@app.route("/login", methods=["GET", "POST"])
def login_page():
    """Trang đăng nhập. Hiện nút Google và/hoặc form username/password."""
    # Nếu chưa cấu hình cơ chế nào thì chuyển thẳng vào app
    if not auth_required():
        return redirect(url_for("index"))

    error = None
    if request.method == "POST":
        # Nhánh username/password (chỉ hoạt động khi đã khai APP_USERNAME)
        if not (AUTH_USERNAME and AUTH_PASSWORD):
            error = "Hệ thống chỉ cho phép đăng nhập bằng Google."
        else:
            username = (request.form.get("username") or "").strip()
            password = request.form.get("password") or ""
            if secrets.compare_digest(username, AUTH_USERNAME) and secrets.compare_digest(password, AUTH_PASSWORD):
                session.clear()
                session["authenticated"] = True
                session.permanent = True
                next_url = request.args.get("next") or url_for("index")
                return redirect(next_url)
            error = "Sai tên đăng nhập hoặc mật khẩu."

    return render_template(
        "login.html",
        error=error,
        google_enabled=auth_module.google_oauth_enabled(),
        password_enabled=bool(AUTH_USERNAME and AUTH_PASSWORD),
    )


@app.route("/login/google")
def login_google():
    """Chuyển hướng sang Google để người dùng chọn tài khoản."""
    if not auth_module.google_oauth_enabled():
        return redirect(url_for("login_page"))
    # ProxyFix đã dựng lại đúng scheme/host từ header của Railway,
    # nên url_for(_external=True) tự sinh https://... chuẩn.
    redirect_uri = url_for("auth_google_callback", _external=True)
    return oauth.google.authorize_redirect(redirect_uri)


@app.route("/auth/google/callback")
def auth_google_callback():
    """Google gọi về sau khi người dùng đồng ý. Tạo phiên đăng nhập tại đây."""
    if not auth_module.google_oauth_enabled():
        return redirect(url_for("login_page"))

    try:
        token = oauth.google.authorize_access_token()
    except Exception as e:
        return render_template("login.html", error=f"Đăng nhập Google thất bại: {e}",
                               google_enabled=True,
                               password_enabled=bool(AUTH_USERNAME and AUTH_PASSWORD))

    info = token.get("userinfo") or {}
    email = (info.get("email") or "").strip().lower()

    if not email:
        return render_template("login.html", error="Không lấy được email từ Google.",
                               google_enabled=True,
                               password_enabled=bool(AUTH_USERNAME and AUTH_PASSWORD))

    if not info.get("email_verified", True):
        return render_template("login.html", error="Email Google chưa được xác minh.",
                               google_enabled=True,
                               password_enabled=bool(AUTH_USERNAME and AUTH_PASSWORD))

    if not auth_module.is_email_allowed(email):
        return render_template("login.html",
                               error=f"Tài khoản {email} không có quyền truy cập công cụ này.",
                               google_enabled=True,
                               password_enabled=bool(AUTH_USERNAME and AUTH_PASSWORD))

    session.clear()
    session["authenticated"] = True
    session["user_email"] = email
    session["user_name"] = info.get("name") or email.split("@")[0]
    session["user_picture"] = info.get("picture") or ""
    session.permanent = True

    auth_module.record_login(DATA_DIR, email,
                             name=session["user_name"],
                             picture=session["user_picture"])

    return redirect(url_for("index"))


@app.route("/api/me", methods=["GET"])
def api_me():
    """Thông tin tài khoản đang đăng nhập — dùng để hiện tên/ảnh trên sidebar."""
    return jsonify({
        "status": "success",
        "data": {
            "authenticated": bool(session.get("authenticated")),
            "email": session.get("user_email", ""),
            "name": session.get("user_name", ""),
            "picture": session.get("user_picture", ""),
            "google_login": auth_module.google_oauth_enabled(),
        }
    })


@app.route("/logout", methods=["GET", "POST"])
def logout_page():
    """Đăng xuất và xóa session."""
    session.clear()
    return redirect(url_for("login_page"))

# Thư mục dữ liệu bền vững.
# Local: dùng thư mục hiện tại.
# Railway: gắn Volume rồi đặt biến DATA_DIR=/data → project/profile không mất khi redeploy.
DATA_DIR = os.environ.get("DATA_DIR", ".")

# --- THƯ MỤC DỮ LIỆU ---
# Khi bật Google OAuth, mỗi tài khoản có thư mục riêng:
#     <DATA_DIR>/users/<hash-email>/projects
#     <DATA_DIR>/users/<hash-email>/profiles
# Hai người dùng khác nhau không bao giờ thấy dữ liệu của nhau.
#
# Khi chạy local không bật OAuth, mọi thứ dồn vào users/local — hành vi
# giống hệt trước đây, không phải đăng nhập vẫn dùng được.

def projects_root() -> str:
    """Thư mục project của NGƯỜI DÙNG HIỆN TẠI."""
    return auth_module.user_projects_root(DATA_DIR)


def profiles_root() -> str:
    """Thư mục profile của NGƯỜI DÙNG HIỆN TẠI."""
    return auth_module.user_profiles_root(DATA_DIR)

# Route chính: Render file index.html từ thư mục templates lên trình duyệt
@app.route('/')
@app.route('/index.html')
def index():
    return render_template('index.html')

# Route cho trang Script Writer độc lập
@app.route('/script_writer.html')
def script_writer_page():
    return render_template('script_writer.html')

# Route cho trang Scene Breakdown (Tool 2)
@app.route('/scene_breakdown.html')
def scene_breakdown_page():
    return render_template('scene_breakdown.html')

# Route cho trang Asset Prompts (Tool 3)
@app.route('/asset_prompts.html')
def asset_prompts_page():
    return render_template('asset_prompts.html')

# Route cho trang Camera Movement (Tool 4)
@app.route('/camera_movement.html')
def camera_movement_page():
    return render_template('camera_movement_v2.html')

# Route cho trang Thumbnail (Tool 5)
@app.route('/thumbnail.html')
def thumbnail_page():
    return render_template('thumbnail_v2.html')

# Route cho trang Video Metadata (Tool 6)
@app.route('/video_metadata.html')
def video_metadata_page():
    return render_template('video_metadata.html')

# Route cho trang Upload QC (Tool 7)
@app.route('/upload_qc.html')
def upload_qc_page():
    return render_template('upload_qc.html')

# --- CÁC API QUẢN LÝ NHIỀU PROJECT (MỖI PROJECT LÀ 1 FILE .JSON TRONG THƯ MỤC CỐ ĐỊNH) ---

@app.route('/api/projects/list', methods=['GET'])
def api_list_projects():
    """Liệt kê tất cả các project dựa trên các file .json có trong thư mục cố định projects/"""
    try:
        files = [f[:-5] for f in os.listdir(projects_root()) if f.endswith('.json')]
        if not files:
            files = ["Default_Project"]
        return jsonify({"status": "success", "data": files})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/projects/load', methods=['POST'])
def api_load_project():
    """Đọc dữ liệu của 1 project từ file .json cố định tương ứng"""
    try:
        data = request.json or {}
        proj_name = data.get("project_name", "Default_Project").strip()
        file_path = os.path.join(projects_root(), f"{proj_name}.json")

        if os.path.exists(file_path):
            with open(file_path, "r", encoding="utf-8") as f:
                proj_data = json.load(f)
            return jsonify({"status": "success", "data": proj_data})
        else:
            return jsonify({"status": "success", "data": {}})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/projects/save', methods=['POST'])
def api_save_project():
    """Lưu dữ liệu của 1 project thành 1 file .json độc lập trong thư mục projects/"""
    try:
        data = request.json or {}
        proj_name = data.get("project_name", "").strip()
        proj_content = data.get("content", {})

        if not proj_name:
            return jsonify({"status": "error", "message": "Tên project không được để trống!"}), 400

        file_path = os.path.join(projects_root(), f"{proj_name}.json")
        with open(file_path, "w", encoding="utf-8") as f:
            json.dump(proj_content, f, ensure_ascii=False, indent=4)

        return jsonify({"status": "success", "message": f"Đã lưu project {proj_name} thành công!"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/projects/delete', methods=['POST'])
def api_delete_project():
    """Xóa file .json của project trong thư mục cố định"""
    try:
        data = request.json or {}
        proj_name = data.get("project_name", "").strip()
        file_path = os.path.join(projects_root(), f"{proj_name}.json")

        if os.path.exists(file_path):
            os.remove(file_path)

        return jsonify({"status": "success", "message": f"Đã xóa project {proj_name}!"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- CÁC API CHO MODULE PROFILES KÊNH (LƯU THÀNH TỪNG FILE ĐỘC LẬP TRONG THƯ MỤC PROFILES/ MANG TÊN KÊNH) ---

@app.route('/api/profiles', methods=['GET'])
def api_get_profiles():
    """Lấy danh sách tất cả các profile kênh từ thư mục cố định profiles/."""
    try:
        profiles = []
        if os.path.exists(profiles_root()):
            for f in os.listdir(profiles_root()):
                if f.endswith('.json'):
                    file_path = os.path.join(profiles_root(), f)
                    try:
                        with open(file_path, "r", encoding="utf-8") as file:
                            data = json.load(file)
                            profiles.append(data)
                    except Exception as e:
                        print(f"Lỗi đọc file profile {f}: {e}")
        return jsonify({"status": "success", "data": profiles})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/profiles', methods=['POST'])
def api_create_profile():
    """Tạo mới và lưu profile thành file .json độc lập có tên đúng bằng tên kênh vào thư mục profiles/."""
    try:
        data = request.json or {}
        ten_kenh = data.get("ten_kenh", "Default_Channel").strip()

        safe_name = re.sub(r'[\\/*?:"<>|]', "", ten_kenh).replace(" ", "_")
        if not safe_name:
            safe_name = "Channel_Profile"

        # ĐÃ SỬA: dùng uuid4 thay vì os.urandom(2) (chỉ có 65536 khả năng, dễ trùng ID khi có nhiều profile)
        profile_id = data.get("id") or uuid.uuid4().hex[:8]
        data["id"] = profile_id

        file_path = os.path.join(profiles_root(), f"{safe_name}.json")

        # ĐÃ SỬA: chống ghi đè nhầm 1 profile KHÁC có cùng tên kênh (safe_name trùng nhưng id khác)
        # -> tránh mất dữ liệu profile cũ một cách âm thầm, và tránh gây lệch ID như lỗi "không tìm thấy profile để xóa"
        if os.path.exists(file_path):
            try:
                with open(file_path, "r", encoding="utf-8") as existing_f:
                    existing_data = json.load(existing_f)
                existing_id = existing_data.get("id")
                if existing_id and str(existing_id) != str(profile_id):
                    suffix = 1
                    while os.path.exists(os.path.join(profiles_root(), f"{safe_name}_{suffix}.json")):
                        suffix += 1
                    file_path = os.path.join(profiles_root(), f"{safe_name}_{suffix}.json")
            except Exception:
                pass

        with open(file_path, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=4)

        return jsonify({"status": "success", "profile_id": profile_id, "message": f"Đã lưu profile [{ten_kenh}] thành công!"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/profiles/<profile_id>', methods=['DELETE'])
def api_delete_profile(profile_id):
    """Xóa file profile tương ứng trong thư mục profiles/ theo ID."""
    try:
        deleted = False
        locked_file = None  # ĐÃ THÊM: theo dõi riêng trường hợp tìm thấy đúng file nhưng bị khóa, không xóa được
        requested_id = str(profile_id).strip()
        seen_ids = []  # để log debug khi không tìm thấy, dễ chẩn đoán hơn lần sau

        if os.path.exists(profiles_root()):
            for f in os.listdir(profiles_root()):
                if f.endswith('.json'):
                    file_path = os.path.join(profiles_root(), f)
                    try:
                        with open(file_path, "r", encoding="utf-8") as file:
                            data = json.load(file)
                            file_id = str(data.get("id", "")).strip()
                            seen_ids.append((f, file_id))
                            if file_id == requested_id:
                                # ĐÃ SỬA: WinError 32 trên Windows thường chỉ là khóa TẠM THỜI do
                                # Windows Defender/dịch vụ đánh index quét file .json vừa ghi/đọc,
                                # không phải do người dùng mở chương trình nào -> tự động thử lại
                                # vài lần (retry) trước khi kết luận là lỗi thật.
                                remove_err = None
                                for attempt in range(5):
                                    try:
                                        os.remove(file_path)
                                        deleted = True
                                        remove_err = None
                                        break
                                    except PermissionError as e:
                                        remove_err = e
                                        time.sleep(0.3 * (attempt + 1))  # đợi tăng dần: 0.3s, 0.6s, 0.9s...

                                if remove_err is not None:
                                    locked_file = f
                                    print(f"[api_delete_profile] Tìm thấy đúng file {f} (id={file_id}) nhưng KHÔNG xóa được sau 5 lần thử vì đang bị khóa: {remove_err}")
                                break
                    except Exception as e:
                        print(f"[api_delete_profile] Lỗi đọc file {f}: {e}")

        if deleted:
            return jsonify({"status": "success", "message": f"Đã xóa thành công profile {profile_id}"})
        elif locked_file:
            # thông báo đúng bản chất lỗi khi file vẫn bị khóa sau khi đã tự động thử lại 5 lần
            return jsonify({
                "status": "error",
                "message": f"File profile ({locked_file}) vẫn đang bị khóa sau nhiều lần thử tự động (có thể do Windows Defender/dịch vụ index đang quét file). Vui lòng thử bấm Xóa lại sau vài giây."
            }), 423
        else:
            # in ra console server toàn bộ id đang có trên đĩa để đối chiếu với id được yêu cầu xóa
            print(f"[api_delete_profile] Không tìm thấy id='{requested_id}'. Các id hiện có trong profiles/: {seen_ids}")
            return jsonify({"status": "error", "message": "Không tìm thấy profile để xóa!"}), 404
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API LƯU GEMINI API KEY ---

@app.route('/api/save-gemini-key', methods=['POST'])
def api_save_gemini_key():
    try:
        data = request.json
        api_key = data.get('api_key')

        if api_key:
            os.environ["GEMINI_API_KEY"] = api_key.strip()
            with open('.env', 'w', encoding='utf-8') as f:
                f.write(f"GEMINI_API_KEY={api_key.strip()}\n")

        return jsonify({"status": "success", "message": "Đã lưu API Key thành công!"})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API CHO AUTO EXTRACT TỪ 3 FILE .MD (ĐÃ KẾT NỐI GEMINI API) ---

@app.route('/api/auto-extract', methods=['POST'])
def api_auto_extract():
    """Nhận nội dung 3 file .md, gọi Gemini API phân tích và trả về các thông số cấu hình."""
    try:
        data = request.json or {}
        style_guide = data.get('style_guide', '')
        dna = data.get('dna', '')
        topic_bank = data.get('topic_bank', '')
        api_key = data.get('api_key', '').strip() or os.getenv("GEMINI_API_KEY", "")

        result = auto_extractor.extract_from_files(style_guide, dna, topic_bank, api_key=api_key)

        return jsonify({"status": "success", "data": result})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API CHO TOPIC IDEATOR ---

@app.route('/api/generate-topics', methods=['POST'])
def api_generate_topics():
    try:
        data = request.json
        result = topic_ideator.process_topics(
            ngach_kenh=data.get('ngach_kenh'),
            so_topics=data.get('so_topics'),
            ngon_ngu=data.get('ngon_ngu'),
            focus=data.get('focus'),
            yeu_cau_bo_sung=data.get('yeu_cau_bo_sung'),
            style_guide=data.get('style_guide', ''),
            dna=data.get('dna', ''),
            topic_bank=data.get('topic_bank', '')
        )
        return jsonify({"status": "success", "data": result})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API EXPORT SCENE LIST THÀNH FILE XLSX ---

@app.route('/api/export-scene-list-xlsx', methods=['POST'])
def api_export_scene_list_xlsx():
    """Xuất Scene List của Tool 2 thành file Excel giống bảng production sheet."""
    try:
        data = request.json or {}
        project_name = data.get("project_name", "scene-list")
        rows = data.get("rows", [])

        if not rows:
            return jsonify({"status": "error", "message": "Không có dữ liệu scene để xuất XLSX!"}), 400

        wb = Workbook()
        ws = wb.active
        ws.title = "Scene List"

        headers = ["#", "Level", "Start", "End", "Duration", "VO", "Character", "Background", "Camera", "Prompt"]
        ws.append(headers)

        header_fill = PatternFill("solid", fgColor="1F4E79")
        header_font = Font(color="FFFFFF", bold=True)
        thin_gray = Side(style="thin", color="D9D9D9")
        border = Border(left=thin_gray, right=thin_gray, top=thin_gray, bottom=thin_gray)

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        for row in rows:
            ws.append([
                row.get("id", ""),
                row.get("level", "Main"),
                row.get("start", ""),
                row.get("end", ""),
                row.get("duration", ""),
                row.get("vo", ""),
                row.get("character", ""),
                row.get("background", ""),
                row.get("camera", ""),
                row.get("prompt", "")
            ])

        widths = [8, 12, 16, 16, 12, 58, 24, 24, 24, 80]
        for index, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(index)].width = width

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        for row_cells in ws.iter_rows(min_row=2):
            for cell in row_cells:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            row_cells[0].alignment = Alignment(horizontal="center", vertical="top")
            row_cells[1].alignment = Alignment(horizontal="center", vertical="top")
            row_cells[2].alignment = Alignment(horizontal="center", vertical="top")
            row_cells[3].alignment = Alignment(horizontal="center", vertical="top")
            row_cells[4].alignment = Alignment(horizontal="center", vertical="top")

        for row_idx in range(2, ws.max_row + 1):
            ws.row_dimensions[row_idx].height = 42

        safe_project_name = re.sub(r'[\\/*?:"<>|]', "", str(project_name)).strip() or "scene-list"
        memory_file = io.BytesIO()
        wb.save(memory_file)
        memory_file.seek(0)

        return send_file(
            memory_file,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=f"{safe_project_name}_scene_list.xlsx"
        )
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API CHO RENAME AND ZIP (SCENE RENAMER) ---

@app.route('/api/rename-and-zip', methods=['POST'])
def api_rename_and_zip():

    """Nhận danh sách file ảnh, đổi tên theo thứ tự 001, 002... và trả về file ZIP"""
    try:
        uploaded_files = request.files.getlist('images')
        start_num = int(request.form.get('start_number', 1))

        if not uploaded_files or uploaded_files[0].filename == '':
            return jsonify({"status": "error", "message": "Chưa chọn file ảnh nào!"}), 400

        memory_file = io.BytesIO()
        with zipfile.ZipFile(memory_file, 'w', zipfile.ZIP_DEFLATED) as zf:
            for index, file_storage in enumerate(uploaded_files):
                ext = os.path.splitext(file_storage.filename)[1].lower()
                if not ext:
                    ext = ".jpg"

                current_num = start_num + index
                new_filename = f"{current_num:03d}{ext}"

                file_bytes = file_storage.read()
                zf.writestr(new_filename, file_bytes)

        memory_file.seek(0)
        return send_file(
            memory_file,
            mimetype='application/zip',
            as_attachment=True,
            download_name='renamed_scenes.zip'
        )
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API CHO MODULE RESEARCH & SCRIPT WRITER ---

@app.route('/api/research_and_write', methods=['POST'])
def api_research_and_write():
    """Nhận thông tin project, đọc 3 file .md gốc, research thực tế qua API và sinh kịch bản."""
    try:
        data = request.json or {}
        topic = data.get("topic", "")
        project_name = data.get("project_name", "")
        profile_id = data.get("profile_id", None)
        options = data.get("options", {})

        api_key = data.get("api_key", "").strip() or os.getenv("GEMINI_API_KEY")

        if not topic:
            return jsonify({"status": "error", "message": "Thiếu chủ đề video (topic)!"}), 400

        if not api_key:
            return jsonify({"status": "error", "message": "Chưa có Gemini API Key! Vui lòng nhập API Key ở sidebar."}), 400

        project_dir = data.get("project_dir", os.path.join(projects_root(), project_name))
        if not os.path.exists(project_dir):
            project_dir = "."

        writer = research_script_writer.ResearchScriptWriter(api_key=api_key)
        result = writer.execute_api_research_and_write(
            topic=topic,
            project_dir=project_dir,
            profile_id=profile_id,
            options=options
        )

        return jsonify(result)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API CHO MODULE DIRECT SCRIPT WRITER (VIẾT NGAY KHÔNG RESEARCH) ---

@app.route('/api/direct_write', methods=['POST'])
def api_direct_write():
    """Nhận thông tin project, viết script nhanh qua Gemini API không sử dụng search."""
    try:
        data = request.json or {}
        topic = data.get("topic", "")
        project_name = data.get("project_name", "")
        profile_id = data.get("profile_id", None)
        options = data.get("options", {})

        api_key = data.get("api_key", "").strip() or os.getenv("GEMINI_API_KEY")

        if not topic:
            return jsonify({"status": "error", "message": "Thiếu chủ đề video (topic)!"}), 400

        if not api_key:
            return jsonify({"status": "error", "message": "Chưa có Gemini API Key! Vui lòng nhập API Key ở sidebar."}), 400

        project_dir = data.get("project_dir", os.path.join(projects_root(), project_name))
        if not os.path.exists(project_dir):
            project_dir = "."

        writer = direct_script_writer.DirectScriptWriter(api_key=api_key)
        result = writer.execute_direct_write(
            topic=topic,
            project_dir=project_dir,
            profile_id=profile_id,
            options=options
        )

        return jsonify(result)
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API CHO MODULE SCENE BREAKDOWN (PRE-SCAN ANALYSIS) ---
# ĐÃ CHUYỂN sang modules/prescan_handler.py để prompt/tool dùng chung ai_client
# và giữ route này chỉ làm controller mỏng.

@app.route('/api/run-prescan', methods=['POST'])
def api_run_prescan():
    """Nhận kịch bản từ Tool 2, gọi AI phân tích và trả về danh sách characters & backgrounds."""
    try:
        data = request.json or {}
        script_text = data.get("script_text", "").strip()
        api_key = data.get("api_key", "").strip() or os.getenv("GEMINI_API_KEY")
        vo_lang = data.get("vo_lang", "Tiếng Việt")
        min_char = data.get("min_char", "30")
        max_char = data.get("max_char", "100")

        result_text = prescan_handler.process_prescan(
            script_text=script_text,
            api_key=api_key,
            vo_lang=vo_lang,
            min_char=min_char,
            max_char=max_char
        )
        return jsonify({"status": "success", "data": result_text})

    except ValueError as e:
        return jsonify({"status": "error", "message": str(e)}), 400
    except RuntimeError as e:
        return jsonify({"status": "error", "message": str(e)}), 500
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API CHO MODULE ASSIGN ASSETS (ĐÃ KẾT NỐI ĐỌC PROMPT STYLE TỪ PROJECT) ---

@app.route('/api/assign-assets', methods=['POST'])
def api_assign_assets():
    """Nhận danh sách phân cảnh, assets, style prompts và project name, tiến hành map và sinh G-Labs Prompts qua API."""
    try:
        data = request.json or {}
        api_key = data.get("api_key", "").strip() or os.getenv("GEMINI_API_KEY")
        scenes_data = data.get("scenes_data", [])
        assets_text = data.get("assets_text", "")
        styles = data.get("styles", {})
        vo_lang = data.get("vo_lang", "Tiếng Việt")

        # Lấy tên project được truyền từ giao diện frontend
        project_name = data.get("project_name", "Default_Project")

        if not api_key:
            return jsonify({"status": "error", "message": "Chưa có API Key!"}), 400

        assigner = assign_assets_handler.AssetAssigner(api_key=api_key)
        # Truyền project_name trực tiếp để AssetAssigner tìm đúng file JSON dự án tương ứng
        result = assigner.process_assignment(
            scenes_data=scenes_data,
            assets_text=assets_text,
            styles=styles,
            vo_lang=vo_lang,
            project_name=project_name
        )

        return jsonify({"status": "success", "data": result})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API CHO TAB "GET-PROMPT IMAGE" ---
# Ghép Scene Style + VO/nhân vật/bối cảnh (đã gán sẵn từ Scene List) thành
# prompt ảnh hoàn chỉnh cho từng scene. Nhân vật/bối cảnh được giữ NGUYÊN VĂN
# y hệt Scene List (xem chi tiết trong modules/get_prompt_image_handler.py).

@app.route('/api/get-prompt-image', methods=['POST'])
def api_get_prompt_image():
    """Nhận scene_style + danh sách scene đã gán nhân vật/bối cảnh, trả về prompt ảnh hoàn chỉnh."""
    try:
        data = request.json or {}
        api_key = data.get("api_key", "").strip() or os.getenv("GEMINI_API_KEY")
        scenes_data = data.get("scenes_data", [])
        scene_style = data.get("scene_style", "")
        # Kết quả Pre-scan: cung cấp mô tả thật cho từng CHARACTERS/BACKGROUNDS ID
        # để AI viết được mô tả bám sát đúng nhân vật & bối cảnh của scene.
        assets_text = data.get("assets_text", "")

        if not api_key:
            return jsonify({"status": "error", "message": "Chưa có API Key!"}), 400

        if not scenes_data:
            return jsonify({"status": "error", "message": "Chưa có dữ liệu Scene List để xử lý!"}), 400

        handler = get_prompt_image_handler.GetPromptImageHandler(api_key=api_key)
        result = handler.process_get_prompt_image(
            scenes_data=scenes_data,
            scene_style=scene_style,
            assets_text=assets_text
        )

        return jsonify({"status": "success", "data": result})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/get-prompt-video', methods=['POST'])
def api_get_prompt_video():
    """
    Tool 2 — Get-Prompt Video.

    Đọc chính kết quả Get-Prompt Image (từng prompt ảnh theo ID scene) cộng với
    VO/nhân vật/bối cảnh/góc máy, rồi viết prompt image-to-video tương ứng cho
    từng scene. Không sinh chuyển động ngẫu nhiên: mỗi prompt video phải bám
    đúng nội dung của prompt ảnh và câu VO của chính scene đó.
    """
    try:
        data = request.json or {}
        api_key = data.get("api_key", "").strip() or os.getenv("GEMINI_API_KEY")
        scenes_data = data.get("scenes_data", [])
        image_prompts_text = data.get("image_prompts_text", "")
        motion_style = data.get("motion_style", "")
        vo_lang = data.get("vo_lang", "Vietnamese")

        if not api_key:
            return jsonify({"status": "error", "message": "Chưa có API Key!"}), 400

        if not scenes_data:
            return jsonify({"status": "error", "message": "Chưa có dữ liệu Scene List để xử lý!"}), 400

        handler = get_prompt_video_handler.GetPromptVideoHandler(api_key=api_key)
        result = handler.process_get_prompt_video(
            scenes_data=scenes_data,
            image_prompts_text=image_prompts_text,
            motion_style=motion_style,
            vo_lang=vo_lang,
        )

        return jsonify({"status": "success", "data": result})
    except ValueError as e:
        return jsonify({"status": "error", "message": str(e)}), 400
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API CHO MODULE ASSET PROMPTS (TOOL 3 — GENERATE ALL PROMPTS) ---

@app.route('/api/generate-asset-prompts', methods=['POST'])
def api_generate_asset_prompts():
    """Nhận char_style/bg_style (prompt dán ngoài) + danh sách characters/backgrounds, sinh bộ prompt hoàn chỉnh."""
    try:
        data = request.json or {}
        api_key = data.get("api_key", "").strip() or os.getenv("GEMINI_API_KEY")
        char_style = data.get("char_style", "")
        bg_style = data.get("bg_style", "")
        scene_style = data.get("scene_style", "")
        characters = data.get("characters", [])
        backgrounds = data.get("backgrounds", [])

        if not api_key:
            return jsonify({"status": "error", "message": "Chưa có API Key! Vui lòng nhập API Key ở sidebar."}), 400

        handler = asset_prompts.AssetPromptsHandler(api_key=api_key)
        result = handler.generate_asset_prompts(
            char_style=char_style,
            bg_style=bg_style,
            scene_style=scene_style,
            characters=characters,
            backgrounds=backgrounds
        )
        return jsonify(result)
    except ValueError as e:
        return jsonify({"status": "error", "message": str(e)}), 400
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API CHO MODULE VIDEO METADATA (TOOL 6 — GENERATE METADATA SEO) ---

@app.route('/api/generate-metadata', methods=['POST'])
def api_generate_metadata():
    """Nhận tiêu đề/kịch bản người dùng dán thủ công + cấu hình SEO, sinh metadata hoàn chỉnh qua AI."""
    try:
        data = request.json or {}
        api_key = data.get("api_key", "").strip() or os.getenv("GEMINI_API_KEY")
        subject = data.get("subject", "")
        summary = data.get("summary", "")
        keyword = data.get("keyword", "")
        cta = data.get("cta", "")
        tone = data.get("tone", "khac")
        output_language = data.get("output_language", "Vietnamese")

        if not api_key:
            return jsonify({"status": "error", "message": "Chưa có API Key! Vui lòng nhập API Key ở sidebar."}), 400

        result = video_metadata_handler.generate_metadata(
            subject=subject,
            summary=summary,
            keyword=keyword,
            cta=cta,
            tone=tone,
            api_key=api_key,
            output_language=output_language
        )
        return jsonify({"status": "success", "data": result})
    except ValueError as e:
        return jsonify({"status": "error", "message": str(e)}), 400
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

# --- API CHO MODULE THUMBNAIL (TOOL 5 — GENERATE THUMBNAIL TITLES & PROMPTS) ---

@app.route('/api/generate-thumbnail-titles', methods=['POST'])
def api_generate_thumbnail_titles():
    """Sinh 5 tiêu đề thumbnail tối ưu CTR từ tiêu đề video HOẶC tóm tắt kịch bản."""
    try:
        data = request.json or {}
        api_key = data.get("api_key", "").strip() or os.getenv("GEMINI_API_KEY")
        video_title = data.get("video_title", "")
        script_summary = data.get("script_summary", "")
        output_lang = data.get("output_lang", "vi")
        model = (data.get("model") or "").strip()

        if not api_key:
            return jsonify({"status": "error", "message": "Chưa có API Key! Vui lòng nhập API Key ở sidebar."}), 400

        result = thumbnail_handler.generate_thumbnail_titles(
            video_title=video_title,
            script_summary=script_summary,
            api_key=api_key,
            output_lang=output_lang,
            model=model
        )
        return jsonify({"status": "success", "data": result})
    except ValueError as e:
        return jsonify({"status": "error", "message": str(e)}), 400
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/analyze-thumbnail-structure', methods=['POST'])
def api_analyze_thumbnail_structure():
    """Phân tích cấu trúc thumbnail: vị trí hình ảnh, fonts chữ, màu sắc & cảm xúc."""
    try:
        data = request.json or {}
        api_key = data.get("api_key", "").strip() or os.getenv("GEMINI_API_KEY")
        video_title = data.get("video_title", "")
        selected_title = data.get("selected_title", "")
        output_lang = data.get("output_lang", "vi")
        model = (data.get("model") or "").strip()

        if not api_key:
            return jsonify({"status": "error", "message": "Chưa có API Key! Vui lòng nhập API Key ở sidebar."}), 400

        result = thumbnail_handler.analyze_thumbnail_structure(
            video_title=video_title,
            selected_title=selected_title,
            api_key=api_key,
            output_lang=output_lang,
            model=model
        )
        return jsonify({"status": "success", "data": result})
    except ValueError as e:
        return jsonify({"status": "error", "message": str(e)}), 400
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500

@app.route('/api/build-thumbnail-prompt', methods=['POST'])
def api_build_thumbnail_prompt():
    """Ghép prompt AI tạo ảnh thumbnail hoàn chỉnh (deterministic, không gọi AI)."""
    try:
        data = request.json or {}
        video_title = data.get("video_title", "")
        selected_topic = data.get("selected_topic", "")
        script_summary = data.get("script_summary", "")
        style = data.get("style", "")
        subtitle = data.get("subtitle", "")
        custom_prompt = data.get("custom_prompt", "")
        image_position = data.get("image_position", "")
        text_style = data.get("text_style", "")
        color_emotion = data.get("color_emotion", "")
        output_lang = data.get("output_lang", "vi")

        result = thumbnail_handler.build_thumbnail_prompt(
            video_title=video_title,
            selected_topic=selected_topic,
            script_summary=script_summary,
            style=style,
            subtitle=subtitle,
            custom_prompt=custom_prompt,
            image_position=image_position,
            text_style=text_style,
            color_emotion=color_emotion,
            output_lang=output_lang
        )
        return jsonify({"status": "success", "data": result})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


@app.route('/api/build-thumbnail-prompt-ai', methods=['POST'])
def api_build_thumbnail_prompt_ai():
    """Tạo prompt thumbnail bằng AI: phân tích ảnh tham chiếu + Active Profile style."""
    try:
        data = request.json or {}
        api_key = data.get("api_key", "").strip() or os.getenv("GEMINI_API_KEY") or ""
        selected_title = data.get("selected_title", "") or data.get("selected_topic", "")
        video_title = data.get("video_title", "")
        script_summary = data.get("script_summary", "")
        subtitle = data.get("subtitle", "")
        custom_prompt = data.get("custom_prompt", "")
        image_position = data.get("image_position", "")
        text_style = data.get("text_style", "")
        color_emotion = data.get("color_emotion", "")
        image_base64 = data.get("image_base64", "")
        output_lang = data.get("output_lang", "vi")
        model = (data.get("model") or "").strip()

        # Resolve profile style: ưu tiên từ frontend gửi lên,
        # fallback đọc trực tiếp từ file profile theo activeProfileId
        char_style = (data.get("char_style") or "").strip()
        bg_style = (data.get("bg_style") or "").strip()
        scene_style = (data.get("scene_style") or "").strip()
        visual = (data.get("visual") or "").strip()
        profile_name = (data.get("profile_name") or "").strip()

        # Nếu chưa có char/bg/scene style (do chưa sync Tool 0), thử resolve qua ID
        active_profile_id = (data.get("active_profile_id") or "").strip()
        if not char_style and not bg_style and not scene_style and not visual and active_profile_id:
            if os.path.exists(profiles_root()):
                for fname in os.listdir(profiles_root()):
                    if not fname.endswith(".json"):
                        continue
                    fpath = os.path.join(profiles_root(), fname)
                    try:
                        with open(fpath, "r", encoding="utf-8") as f:
                            pj = json.load(f)
                        if str(pj.get("id", "")) == str(active_profile_id):
                            char_style = char_style or (pj.get("char_style") or "")
                            bg_style = bg_style or (pj.get("bg_style") or "")
                            scene_style = scene_style or (pj.get("scene_style") or "")
                            visual = visual or (pj.get("visual") or "")
                            profile_name = profile_name or (pj.get("ten_kenh") or "")
                            break
                    except Exception:
                        continue

        result = thumbnail_handler.build_thumbnail_prompt_ai(
            api_key=api_key,
            selected_title=selected_title,
            video_title=video_title,
            script_summary=script_summary,
            subtitle=subtitle,
            custom_prompt=custom_prompt,
            image_position=image_position,
            text_style=text_style,
            color_emotion=color_emotion,
            image_base64=image_base64,
            char_style=char_style,
            bg_style=bg_style,
            scene_style=scene_style,
            visual=visual,
            profile_name=profile_name,
            output_lang=output_lang,
            model=model,
        )
        return jsonify({"status": "success", "data": result})
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


# --- API CHO TOOL 4 (CAMERA & SHOT DESIGNER) ---

@app.route('/api/analyze-camera', methods=['POST'])
def api_analyze_camera():
    """
    Tool 4 — Phân tích SRT hoặc scene list từ Tool 2, thiết kế góc máy bằng AI.

    Input (JSON):
        - srt_content: nội dung file .srt (text thuần)
        - scenes: hoặc danh sách scene từ Tool 2 (list of dict với timecode, text, character, background)
        - api_key: Gemini API key hoặc ShopAIKey
        - genre: "documentary" | "tech_report" | "drama" | "explainer"
        - transition_default: "Smooth Cut" | "Hard Cut" | "Fade" | "Match Cut"
        - vfx_style: "Cinematic Depth" | "Flat / Clean" | "High Contrast" | "Motion Graphic"
        - max_shot_seconds: ngưỡng giây (mặc định 6)
        - constraints: {split_long, no_repeat, open_wide, axis_180}
        - style_guide: chuỗi style guide từ active profile (tùy chọn)
        - model: tên model AI (tùy chọn, mặc định gemini-2.5-flash)

    Output:
        {
          "status": "success",
          "data": {
            "shots": [...],      // danh sách shot đã chuẩn hóa
            "metrics": {...},    // 6 chỉ số cho stat strip
            "warnings": [...],   // cảnh báo dựng
            "source": "ai" | "fallback",
            "note": "..."        // lý do dùng fallback (nếu có)
          }
        }
    """
    try:
        data = request.json or {}

        srt_content = data.get("srt_content", "")
        scenes = data.get("scenes") or []
        api_key = (data.get("api_key") or "").strip() or os.getenv("GEMINI_API_KEY") or ""

        genre = data.get("genre", "explainer")
        transition_default = data.get("transition_default", "Smooth Cut")
        vfx_style = data.get("vfx_style", "Cinematic Depth")
        max_shot_seconds = float(data.get("max_shot_seconds") or 6)
        constraints = data.get("constraints") or {}
        style_guide = data.get("style_guide", "")
        model = (data.get("model") or "").strip()

        # Gọi handler
        result = camera_handler.camera_handler.analyze(
            srt_content=srt_content,
            scenes=scenes,
            api_key=api_key,
            genre=genre,
            transition_default=transition_default,
            vfx_style=vfx_style,
            max_shot_seconds=max_shot_seconds,
            constraints=constraints,
            style_guide=style_guide,
            model=model
        )

        return jsonify({"status": "success", "data": result})

    except ValueError as e:
        return jsonify({"status": "error", "message": str(e)}), 400
    except Exception as e:
        return jsonify({"status": "error", "message": str(e)}), 500


if __name__ == '__main__':
    # Local: debug=True. Trên Railway/production: dùng gunicorn (Procfile), không chạy khối này.
    port = int(os.environ.get("PORT", 5000))
    debug = os.environ.get("FLASK_DEBUG", "1") == "1"
    app.run(host="0.0.0.0", debug=debug, port=port)